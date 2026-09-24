import docx
from pptx import Presentation
import openpyxl

def extract_docx(file_path: str) -> str:
    """Extracts text from a Word document."""
    try:
        doc = docx.Document(file_path)
        full_text = []
        for para in doc.paragraphs:
            if para.text.strip():
                full_text.append(para.text)
        return "\n".join(full_text)
    except Exception as e:
        return f"[Error parsing DOCX: {str(e)}]"

def extract_pptx(file_path: str) -> str:
    """Extracts text from a PowerPoint presentation."""
    try:
        prs = Presentation(file_path)
        full_text = []
        for slide_num, slide in enumerate(prs.slides, 1):
            slide_text = []
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    if shape.text.strip():
                        slide_text.append(shape.text)
            if slide_text:
                full_text.append(f"--- Slide {slide_num} ---\n" + "\n".join(slide_text))
        return "\n\n".join(full_text)
    except Exception as e:
        return f"[Error parsing PPTX: {str(e)}]"

def extract_xlsx(file_path: str) -> str:
    """Extracts text from an Excel spreadsheet."""
    try:
        wb = openpyxl.load_workbook(file_path, data_only=True)
        full_text = []
        for sheet_name in wb.sheetnames:
            sheet = wb[sheet_name]
            sheet_data = []
            for row in sheet.iter_rows(values_only=True):
                # Filter out None values and convert to string
                row_values = [str(cell) for cell in row if cell is not None]
                if row_values:
                    sheet_data.append(" | ".join(row_values))
            if sheet_data:
                full_text.append(f"--- Sheet: {sheet_name} ---\n" + "\n".join(sheet_data))
        return "\n\n".join(full_text)
    except Exception as e:
        return f"[Error parsing XLSX: {str(e)}]"
